#!/usr/bin/env python3
"""
Build a QA test-case workbook with two sheets, matching the team's template:

  "E2E Test Cases" sheet:
    No | Test Case Name | Module 1..5 | Priority | Precondition
    | Test Steps / Prompt | Expected Result | Result (Desktop | Android | iOS) | Remarks

  "API Test Cases" sheet:
    No | Test Case Name | Module 1..5 | Priority | Precondition | Method | Endpoint
    | Request Payload | Test Steps / Prompt | Expected Status Code | Expected Result
    | Result | Remarks

Usage:
    python3 generate_xlsx.py --input cases.json --output test_cases.xlsx

Input JSON shape:
{
  "e2e": [
    {
      "no": 1,
      "test_case_name": "Open Capability list via @ in composer",
      "modules": ["Dropdown - First-Level"],        # 1-5 strings, top-down breakdown
      "priority": "P0",                              # P0 / P1 / P2
      "precondition": "Composer is empty and focused.",
      "steps": "1. Click composer\n2. Type '@'",
      "expected_result": "- Combined dropdown opens...\n- Highlights ... by default",
      "result": {"desktop": "", "android": "", "ios": ""},  # leave blank if not executed
      "remarks": ""
    }
  ],
  "api": [
    {
      "no": 1,
      "test_case_name": "Reject order creation with negative quantity",
      "modules": ["Orders", "Create"],
      "priority": "P0",
      "precondition": "Valid auth token, existing product SKU.",
      "method": "POST",
      "endpoint": "/api/v1/orders",
      "request_payload": "{\\"sku\\": \\"ABC123\\", \\"quantity\\": -1}",
      "steps": "1. Send POST /api/v1/orders with negative quantity",
      "expected_status_code": "400",
      "expected_result": "Response body contains error code INVALID_QUANTITY; no order created.",
      "result": "",
      "remarks": ""
    }
  ]
}

Either list may be omitted/empty if that test type doesn't apply, but at
least one of "e2e" / "api" must contain cases.
"""
import argparse
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit(
        "openpyxl is required but not installed.\n"
        "Run: pip3 install --quiet openpyxl\nThen re-run this script."
    )

# Each column is (label, subcols). subcols is None for a plain column that
# merges vertically across the two header rows; a tuple of sub-labels means
# the column merges horizontally on row 1 and shows the sub-labels on row 2.
E2E_COLUMNS = [
    ("No", None),
    ("Test Case Name", None),
    ("Module 1", None),
    ("Module 2", None),
    ("Module 3", None),
    ("Module 4", None),
    ("Module 5", None),
    ("Priority", None),
    ("Precondition", None),
    ("Test Steps / Prompt", None),
    ("Expected Result", None),
    ("Result", ("Desktop", "Android", "iOS")),
    ("Remarks", None),
]

API_COLUMNS = [
    ("No", None),
    ("Test Case Name", None),
    ("Module 1", None),
    ("Module 2", None),
    ("Module 3", None),
    ("Module 4", None),
    ("Module 5", None),
    ("Priority", None),
    ("Precondition", None),
    ("Method", None),
    ("Endpoint", None),
    ("Request Payload", None),
    ("Test Steps / Prompt", None),
    ("Expected Status Code", None),
    ("Expected Result", None),
    ("Result", None),
    ("Remarks", None),
]

WRAP_COLUMNS = {
    "Test Case Name",
    "Precondition",
    "Request Payload",
    "Test Steps / Prompt",
    "Expected Result",
    "Remarks",
}
COLUMN_WIDTHS = {
    "No": 5,
    "Test Case Name": 32,
    "Module 1": 16,
    "Module 2": 16,
    "Module 3": 16,
    "Module 4": 16,
    "Module 5": 16,
    "Priority": 9,
    "Precondition": 28,
    "Method": 9,
    "Endpoint": 26,
    "Request Payload": 34,
    "Test Steps / Prompt": 38,
    "Expected Status Code": 12,
    "Expected Result": 42,
    "Desktop": 10,
    "Android": 10,
    "iOS": 10,
    "Result": 12,
    "Remarks": 20,
}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def flat_headers(columns_spec):
    flat = []
    for label, subcols in columns_spec:
        flat.extend(subcols if subcols else [label])
    return flat


def write_headers(ws, columns_spec):
    col = 1
    for label, subcols in columns_spec:
        if subcols:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(subcols) - 1)
            for i, sub in enumerate(subcols):
                c = ws.cell(row=2, column=col + i, value=sub)
                c.font, c.fill, c.alignment, c.border = HEADER_FONT, HEADER_FILL, CENTER, BORDER
            top_cell = ws.cell(row=1, column=col, value=label)
            col += len(subcols)
        else:
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            top_cell = ws.cell(row=1, column=col, value=label)
            col += 1
        top_cell.font, top_cell.fill, top_cell.alignment, top_cell.border = HEADER_FONT, HEADER_FILL, CENTER, BORDER


def write_rows(ws, headers, rows):
    r = 3
    for row_values in rows:
        for c_idx, value in enumerate(row_values, start=1):
            label = headers[c_idx - 1]
            cell = ws.cell(row=r, column=c_idx, value=value)
            cell.border = BORDER
            cell.alignment = LEFT_WRAP if label in WRAP_COLUMNS else CENTER
        r += 1


def padded_modules(case):
    modules = (case.get("modules") or []) + [""] * 5
    return modules[:5]


def e2e_row(case):
    result = case.get("result") or {}
    return (
        [case.get("no", ""), case.get("test_case_name", "")]
        + padded_modules(case)
        + [case.get("priority", ""), case.get("precondition", "")]
        + [case.get("steps", ""), case.get("expected_result", "")]
        + [result.get("desktop", ""), result.get("android", ""), result.get("ios", "")]
        + [case.get("remarks", "")]
    )


def api_row(case):
    return (
        [case.get("no", ""), case.get("test_case_name", "")]
        + padded_modules(case)
        + [case.get("priority", ""), case.get("precondition", "")]
        + [case.get("method", ""), case.get("endpoint", ""), case.get("request_payload", "")]
        + [case.get("steps", ""), case.get("expected_status_code", ""), case.get("expected_result", "")]
        + [case.get("result", ""), case.get("remarks", "")]
    )


def add_sheet(wb, ws, title, columns_spec, cases, row_builder):
    ws.title = title[:31]
    write_headers(ws, columns_spec)
    headers = flat_headers(columns_spec)
    for idx, label in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = COLUMN_WIDTHS.get(label, 16)
    write_rows(ws, headers, [row_builder(c) for c in cases])
    ws.freeze_panes = "C3"


def build_workbook(e2e_cases, api_cases):
    if not e2e_cases and not api_cases:
        sys.exit("No test cases provided — both 'e2e' and 'api' lists are empty.")

    wb = Workbook()
    used_default = False

    if e2e_cases:
        add_sheet(wb, wb.active, "E2E Test Cases", E2E_COLUMNS, e2e_cases, e2e_row)
        used_default = True

    if api_cases:
        ws = wb.active if not used_default else wb.create_sheet()
        add_sheet(wb, ws, "API Test Cases", API_COLUMNS, api_cases, api_row)

    return wb


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True, help='Path to JSON file: {"e2e": [...], "api": [...]}')
    ap.add_argument("--output", required=True, help="Path to write the .xlsx file")
    args = ap.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        sys.exit('Input JSON must be an object like {"e2e": [...], "api": [...]}.')

    e2e_cases = data.get("e2e") or []
    api_cases = data.get("api") or []

    wb = build_workbook(e2e_cases, api_cases)
    wb.save(args.output)
    print(f"Wrote {len(e2e_cases)} E2E case(s) and {len(api_cases)} API case(s) to {args.output}")


if __name__ == "__main__":
    main()
